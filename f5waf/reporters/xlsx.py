"""Excel raporu: Özet, Cihazlar, Partition Risk, Bulgular, kontrol başına sheet, Kapanan, İstisnalar."""
from __future__ import annotations

import re

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import SEVERITIES

SEV_COLORS = {"CRITICAL": "B3132B", "HIGH": "E0531F", "MEDIUM": "E0A800", "LOW": "2F7FD1", "INFO": "8A929C"}
HDR_FILL = PatternFill("solid", fgColor="1F2A44")
HDR_FONT = Font(color="FFFFFF", bold=True)

FINDING_COLS = [("severity", "Önem", 11), ("status", "Durum", 12), ("check_id", "Kontrol", 10),
                ("title", "Başlık", 38), ("device", "Cihaz", 14), ("partition", "Partition", 18),
                ("customer", "Müşteri", 22), ("object_type", "Nesne tipi", 13), ("object_name", "Nesne", 42),
                ("detail", "Detay", 60), ("recommendation", "Öneri", 60), ("first_seen", "İlk görülme", 20)]


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", title)[:31])
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _color_sev(ws, col_letter="A"):
    rng = f"{col_letter}2:{col_letter}{max(2, ws.max_row)}"
    for s, color in SEV_COLORS.items():
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{s}"'], fill=PatternFill("solid", fgColor=color),
            font=Font(color="FFFFFF", bold=True)))


def _finding_rows(findings):
    return [[f.get(k, "") for k, _, _ in FINDING_COLS] for f in findings]


def write_xlsx(report: dict, path: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    s = report["summary"]
    active = [f for f in report["findings"] if not f["suppressed"]]

    ws = _sheet(wb, "Özet", ["Alan", "Değer"], [
        ["Tarama ID", report["scan_id"]], ["Başlangıç", report["started_at"]], ["Bitiş", report["finished_at"]],
        ["Durum", report["status"]], ["Taranan cihaz", f"{s['devices_ok']}/{s['devices_total']}"],
        ["Partition sayısı", len(s["partitions"])], ["Toplam aktif bulgu", s["total"]],
        *[[sev, s["by_severity"][sev]] for sev in SEVERITIES],
        ["Yeni bulgu", s["new"]], ["Kapanan bulgu", s["resolved"]], ["İstisnalı bulgu", s["suppressed"]],
    ], [26, 40])
    ws.append([])
    ws.append(["Kontrol", "Bulgu sayısı"])
    for c in s["by_check"]:
        ws.append([f"{c['check_id']} · {c['title']}", c["count"]])

    _sheet(wb, "Cihazlar", ["Cihaz", "Host", "Sürüm", "Durum", "Hata", "Partition", "VS", "Pool", "WAF policy", "Süre (sn)"],
           [[d["name"], d["host"], d["version"], "OK" if d["ok"] else "HATA", d["error"], len(d["partitions"]),
             d["stats"].get("virtuals"), d["stats"].get("pools"), d["stats"].get("policies"), d["duration_sec"]]
            for d in report["devices"]], [14, 26, 12, 8, 50, 10, 8, 8, 11, 10])

    ws = _sheet(wb, "Partition Risk", ["Cihaz", "Partition", "Müşteri", "Risk skoru", *SEVERITIES],
                [[p["device"], p["partition"], p["customer"], p["score"], *[p[x] for x in SEVERITIES]]
                 for p in s["partitions"]], [14, 22, 28, 11, 10, 10, 10, 10, 10])
    for i, sev in enumerate(SEVERITIES[:4], start=5):
        col = get_column_letter(i)
        ws.conditional_formatting.add(f"{col}2:{col}{max(2, ws.max_row)}", CellIsRule(
            operator="greaterThan", formula=["0"], font=Font(color=SEV_COLORS[sev], bold=True)))

    heads, widths = [h for _, h, _ in FINDING_COLS], [w for _, _, w in FINDING_COLS]
    ws = _sheet(wb, "Tüm Bulgular", heads, _finding_rows(active), widths)
    _color_sev(ws)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    by_check: dict[str, list] = {}
    for f in active:
        by_check.setdefault(f["check_id"], []).append(f)
    for cid in sorted(by_check):
        ws = _sheet(wb, f"{cid}", heads, _finding_rows(by_check[cid]), widths)
        _color_sev(ws)

    ws = _sheet(wb, "Kapanan", heads, _finding_rows(report.get("resolved", [])), widths)
    _color_sev(ws)
    supp = [f for f in report["findings"] if f["suppressed"]]
    _sheet(wb, "İstisnalar", heads + ["İstisna nedeni"],
           [r + [f["suppress_reason"]] for r, f in zip(_finding_rows(supp), supp)], widths + [40])
    _sheet(wb, "Kontrol Kataloğu", ["ID", "Başlık", "Kategori", "Açıklama", "Öneri"],
           [[c["id"], c["title"], c["category"], c["description"], c["recommendation"]] for c in report["catalog"]],
           [10, 40, 16, 70, 70])
    wb.save(path)
    return path
