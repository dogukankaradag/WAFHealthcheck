"""Birim + entegrasyon testleri.  Çalıştırma:  python -m pytest -q"""
from __future__ import annotations

import copy
import os
import sys
from datetime import datetime, timedelta, timezone

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from f5waf.checks import run_checks  # noqa: E402
from f5waf.client import F5Client  # noqa: E402
from f5waf.collector import DeviceSnapshot, PolicyData  # noqa: E402
from f5waf.models import Config  # noqa: E402
from f5waf.scanner import Scanner  # noqa: E402
from f5waf.state import StateStore  # noqa: E402

NOW = datetime.now(timezone.utc)


# ----------------------------------------------------------------- client
class FakeResp:
    def __init__(self, code, data):
        self.status_code, self._d, self.text = code, data, str(data)

    def json(self):
        return self._d


class FakeSession:
    def __init__(self, routes):
        self.routes, self.headers, self.verify, self.calls = routes, {}, True, []

    def post(self, url, json=None, timeout=None):
        return FakeResp(200, {"token": {"token": "T1"}})

    def patch(self, url, json=None, timeout=None):
        return FakeResp(200, {})

    def delete(self, url, timeout=None):
        return FakeResp(200, {})

    def get(self, url, params=None, timeout=None):
        self.calls.append(url)
        r = self.routes[url]
        return r.pop(0) if isinstance(r, list) else r


def test_client_pagination_rewrites_localhost_and_relogins():
    c = F5Client("bigip.lab", "u", "p")
    base = "https://bigip.lab/mgmt/tm/asm/policies"
    c.session = FakeSession({
        base: [FakeResp(401, {}), FakeResp(200, {"items": [{"id": 1}], "nextLink": "https://localhost/mgmt/tm/asm/policies?$skip=1"})],
        f"{base}?$skip=1": FakeResp(200, {"items": [{"id": 2}]}),
    })
    items = c.get_items("/mgmt/tm/asm/policies")
    assert [i["id"] for i in items] == [1, 2]
    assert c.session.headers["X-F5-Auth-Token"] == "T1"


def test_client_stats_flatten():
    c = F5Client("bigip.lab", "u", "p")
    c.session = FakeSession({"https://bigip.lab/mgmt/tm/ltm/virtual/stats": FakeResp(200, {"entries": {
        "https://localhost/mgmt/tm/ltm/virtual/~P_A~vs1/stats": {"nestedStats": {"entries": {
            "tmName": {"description": "/P_A/vs1"},
            "status.availabilityState": {"description": "offline"},
            "clientside.curConns": {"value": 3}}}}}})})
    st = c.get_stats("/mgmt/tm/ltm/virtual/stats")
    assert st["/P_A/vs1"]["status.availabilityState"] == "offline"
    assert st["/P_A/vs1"]["clientside.curConns"] == 3


# ----------------------------------------------------------------- checks
def _policy(mode="blocking", days_since_change=1, vs=("/P_A/vs1",), **kw):
    raw = {"id": "x", "name": "pol", "partition": "P_A", "fullPath": "/P_A/pol", "enforcementMode": mode,
           "active": True, "virtualServers": list(vs),
           "createdDatetime": (NOW - timedelta(days=400)).isoformat(),
           "versionDatetime": (NOW - timedelta(days=days_since_change)).isoformat()}
    pd = PolicyData(raw=raw, general={"enforcementReadinessPeriod": 7}, policy_builder={"learningMode": "manual"},
                    staged_signatures=0, disabled_signatures=0)
    for k, v in kw.items():
        setattr(pd, k, v)
    return pd


def _snap(policies):
    s = DeviceSnapshot("waf01", "h", NOW, partitions=["P_A"])
    s.virtuals = [{"fullPath": "/P_A/vs1", "partition": "P_A", "pool": "/P_A/pool1",
                   "profilesReference": {"items": [{"fullPath": "/Common/http"}]},
                   "securityLogProfiles": ["log"]}]
    s.virtual_stats = {"/P_A/vs1": {"status.availabilityState": "available", "status.enabledState": "enabled"}}
    s.pools = [{"fullPath": "/P_A/pool1", "partition": "P_A", "monitor": "/Common/http",
                "membersReference": {"items": [{"name": "a", "state": "up"}]}}]
    s.http_profiles = {"/Common/http"}
    s.policies = policies
    s.signature_update = NOW - timedelta(days=3)
    s.sync_status = "In Sync"
    return s


def _run(snap, cfg=None, observe=None):
    return run_checks(snap, cfg or Config({}), observe or (lambda k: NOW))


@pytest.mark.parametrize("days,sev", [(3, "MEDIUM"), (20, "HIGH"), (90, "CRITICAL")])
def test_transparent_severity_by_age(days, sev):
    f = [x for x in _run(_snap([_policy("transparent", days)])) if x.check_id == "WAF-001"]
    assert len(f) == 1 and f[0].severity == sev


def test_transparent_uses_observed_history_when_longer():
    # cihazda 2 gün önce değişiklik olmuş ama state DB 30 gündür transparent görüyor
    f = _run(_snap([_policy("transparent", 2)]), observe=lambda k: NOW - timedelta(days=30))
    assert [x.severity for x in f if x.check_id == "WAF-001"] == ["HIGH"]


def test_partition_override_threshold():
    cfg = Config({"partition_overrides": {"P_*": {"transparent_max_days": 45}}})
    f = _run(_snap([_policy("transparent", 20)]), cfg)
    assert [x.severity for x in f if x.check_id == "WAF-001"] == ["MEDIUM"]


def test_clean_policy_has_no_findings():
    assert _run(_snap([_policy()])) == []


def test_signature_set_and_violation_block():
    p = _policy(signature_sets=[{"signatureSetReference": {"name": "SQLi"}, "alarm": True, "block": False}],
                violations=[{"description": "Attack signature detected", "alarm": True, "block": False},
                            {"description": "Illegal URL", "alarm": False, "block": False}])  # kritik listede değil
    f = _run(_snap([p]))
    assert {(x.check_id, x.severity) for x in f} == {("WAF-005", "HIGH"), ("WAF-006", "HIGH")}


def test_staging_old_policy():
    f = _run(_snap([_policy(staged_signatures=40)]))
    assert [x.check_id for x in f] == ["WAF-004"]


def test_http_vs_without_waf_and_no_logging():
    s = _snap([])
    assert [x.check_id for x in _run(s)] == ["LTM-003"]
    s2 = _snap([_policy()])
    s2.virtuals[0]["securityLogProfiles"] = []
    assert [x.check_id for x in _run(s2)] == ["LTM-004"]


@pytest.mark.parametrize("dest,port", [("/P/10.0.0.1:443", 443), ("/P/10.0.0.1%12:80", 80),
                                       ("/P/2001:db8::1.8443", 8443), ("/P/0.0.0.0:any", 0), (None, None)])
def test_vs_port_parse(dest, port):
    from f5waf.checks import vs_port
    assert vs_port(dest) == port


def test_no_waf_vs_excludes_port_80_and_lists_l4():
    s = _snap([_policy()])  # vs1 WAF'lı
    s.virtuals += [
        {"fullPath": "/P_A/vs_80", "partition": "P_A", "destination": "/P_A/10.0.0.1:80",
         "rules": ["/Common/_sys_https_redirect"], "profilesReference": {"items": [{"fullPath": "/Common/http"}]}},
        {"fullPath": "/P_A/vs_8080", "partition": "P_A", "destination": "/P_A/10.0.0.1:8080", "pool": "/P_A/pool1",
         "profilesReference": {"items": [{"fullPath": "/Common/http"}]}},
        {"fullPath": "/P_A/vs_l4", "partition": "P_A", "destination": "/P_A/10.0.0.1:443", "pool": "/P_A/pool1",
         "profilesReference": {"items": [{"fullPath": "/Common/fastL4"}]}}]
    s.virtual_stats["/P_A/vs_80"] = {"status.availabilityState": "unknown"}
    got = {(x.object_name, x.severity) for x in _run(s) if x.check_id in ("LTM-003", "LTM-001")}
    assert got == {("/P_A/vs_8080", "HIGH"), ("/P_A/vs_l4", "MEDIUM")}
    cfg = Config({"scan": {"no_waf_exempt_ports": [], "no_waf_include_non_http": False}})
    got = {x.object_name for x in _run(s, cfg) if x.check_id == "LTM-003"}
    assert got == {"/P_A/vs_80", "/P_A/vs_8080"}


def test_pool_all_down_and_unreadable_policy():
    s = _snap([_policy(errors=["violations: 503"])])
    s.pools[0]["membersReference"]["items"][0]["state"] = "down"
    ids = sorted((x.check_id, x.severity) for x in _run(s))
    assert ids == [("LTM-002", "HIGH"), ("WAF-011", "MEDIUM")]


def test_suppression_and_expiry():
    cfg = Config({"suppressions": [
        {"check": "WAF-001", "partition": "P_A", "reason": "CHG1", "until": "2999-01-01"},
        {"check": "LTM-003", "reason": "eski", "until": "2000-01-01"}]})
    from f5waf.models import Finding
    f1 = Finding("WAF-001", "", "HIGH", "", "waf01", "P_A", "", "/P_A/pol", "", "")
    f2 = Finding("LTM-003", "", "HIGH", "", "waf01", "P_A", "", "/P_A/vs", "", "")
    assert cfg.match_suppression(f1).startswith("CHG1")
    assert cfg.match_suppression(f2) is None


# ------------------------------------------------------- scanner + state
def _mock_cfg(tmp_path):
    return Config({"devices": [{"name": n, "mock": True} for n in ("waf01", "waf02", "waf03")],
                   "scan": {"exclude_partitions": ["Common"]},
                   "state_db": str(tmp_path / "s.db"), "output": {"dir": str(tmp_path / "r")}})


def test_scan_tracks_new_persistent_resolved(tmp_path, monkeypatch):
    cfg = _mock_cfg(tmp_path)
    store = StateStore(cfg.state_db)
    r1 = Scanner(cfg, store).run()
    assert r1.status == "completed" and r1.findings and all(f.status == "NEW" for f in r1.findings)

    # ikinci taramada bir policy'yi düzelt -> kapanan bulgu olmalı
    from f5waf import mock
    orig = mock.MockClient.__init__

    def patched(self, name, profile):
        orig(self, name, profile)
        for p in self.policies:
            p["enforcementMode"] = "blocking"
    monkeypatch.setattr(mock.MockClient, "__init__", patched)
    r2 = Scanner(cfg, store).run()
    assert not [f for f in r2.findings if f.check_id == "WAF-001"]
    assert any(f["check_id"] == "WAF-001" for f in r2.resolved)
    assert all(f.status == "PERSISTENT" for f in r2.findings if f.check_id == "CERT-001")


def test_partial_scan_scope(tmp_path):
    cfg = _mock_cfg(tmp_path)
    store = StateStore(cfg.state_db)
    Scanner(cfg, store).run()
    r = Scanner(cfg, store).run(devices=["waf01"], partitions=["P_ACME"])
    assert {f.partition for f in r.findings} <= {"P_ACME", "(cihaz)"}
    assert r.resolved == []  # taranmayan partition'lar "kapandı" sayılmamalı


def test_reports_written(tmp_path):
    from f5waf.reporters import write_all
    cfg = _mock_cfg(tmp_path)
    rep = Scanner(cfg).run().to_report()
    files = write_all(rep, str(tmp_path / "out"), ["html", "xlsx", "json", "csv"])
    for p in files.values():
        assert os.path.getsize(p) > 1000


# -------------------------------------------------------------------- api
@pytest.fixture()
def client(tmp_path, monkeypatch):
    import importlib
    import time as _t

    import yaml
    cfgp = tmp_path / "c.yaml"
    cfgp.write_text(yaml.safe_dump({"devices": [{"name": "waf01", "mock": True}, {"name": "waf02", "mock": True}],
                                    "scan": {"exclude_partitions": ["Common"]},
                                    "state_db": str(tmp_path / "a.db"), "output": {"dir": str(tmp_path / "r")},
                                    "api": {"keys": ["k1"]}}))
    monkeypatch.setenv("F5WAF_CONFIG", str(cfgp))
    monkeypatch.setenv("F5WAF_ADMIN_PASSWORD", "Admin-Parola-1")
    import f5waf.api as api
    importlib.reload(api)
    from fastapi.testclient import TestClient
    with TestClient(api.app) as c:
        c.wait_scan = lambda: next(None for _ in range(200)
                                   if not c.get("/api/scan/status", headers=c.H).json()["running"] or _t.sleep(0.05))
        yield c


def _login(c, user="admin", pw="Admin-Parola-1"):
    r = c.post("/api/auth/login", data={"username": user, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_api_auth_and_roles(client):
    c = client
    assert c.get("/api/summary").status_code == 401
    assert c.post("/api/auth/login", data={"username": "admin", "password": "yanlis"}).status_code == 401
    H = c.H = _login(c)
    assert c.get("/api/auth/me", headers=H).json()["role"] == "admin"
    # readonly kullanıcı oluştur
    assert c.post("/api/auth/users", json={"username": "ro", "password": "kisa"}, headers=H).status_code == 400
    assert c.post("/api/auth/users", json={"username": "ro", "password": "Okuma-1234"}, headers=H).status_code == 201
    RO = _login(c, "ro", "Okuma-1234")
    assert c.post("/api/scan", json={}, headers=RO).status_code == 403
    assert c.get("/api/auth/users", headers=RO).status_code == 403
    # son admin silinemez / pasif yapılamaz
    me = c.get("/api/auth/me", headers=H).json()
    assert c.delete(f"/api/auth/users/{me['id']}", headers=H).status_code == 400
    # API anahtarı admin sayılır
    assert c.get("/api/auth/users", headers={"X-API-Key": "k1"}).status_code == 200
    # brute force kilidi
    for _ in range(5):
        c.post("/api/auth/login", data={"username": "ro", "password": "x"})
    assert c.post("/api/auth/login", data={"username": "ro", "password": "Okuma-1234"}).status_code == 429


def test_api_scan_flow(client):
    c = client
    H = c.H = _login(c)
    s = c.get("/api/summary", headers=H).json()
    assert s["session_id"] is None
    assert c.post("/api/scan", json={"device_ids": ["yok"]}, headers=H).status_code == 400
    r = c.post("/api/scan", json={"device_ids": ["waf01"], "partitions": "P_ACME, P_D*"}, headers=H).json()
    sid = r["session_id"]
    c.wait_scan()
    st = c.get("/api/scan/status", headers=H).json()
    assert st["scan"]["id"] == sid and st["scan"]["status"] == "completed"
    assert st["progress"]["waf01"]["state"] == "done"
    hist = c.get("/api/scan/history", headers=H).json()
    assert hist[0]["triggered_by"] == "admin" and hist[0]["request"]["partitions"] == ["P_ACME", "P_D*"]
    devs = c.get("/api/devices", headers=H).json()
    assert [d["id"] for d in devs] == ["waf01"]
    assert {p["name"] for p in devs[0]["partitions"]} <= {"P_ACME", "P_DELTA_LOG", "(cihaz)"}
    f = c.get("/api/findings", params={"severity": "high"}, headers=H).json()
    assert f["total"] > 0 and all(x["severity"] == "HIGH" for x in f["findings"])
    assert c.get("/api/partitions", headers=H).json()[0]["score"] > 0
    assert c.get("/api/findings/check-names", headers=H).json()

    # iş akışı: çözüldü -> açık listeden düşer, özet sayacı artar
    total0 = c.get("/api/summary", headers=H).json()["total_findings"]
    fp = f["findings"][0]["fingerprint"]
    assert c.patch(f"/api/findings/{fp}/status", json={"status": "bad"}, headers=H).status_code == 400
    assert c.patch(f"/api/findings/{fp}/status", json={"status": "resolved", "comment": "CHG1"},
                   headers=H).status_code == 200
    s = c.get("/api/summary", headers=H).json()
    assert s["total_findings"] == total0 - 1 and s["resolved_count"] == 1
    res = c.get("/api/findings", params={"status": "resolved"}, headers=H).json()["findings"]
    assert res[0]["fingerprint"] == fp and res[0]["status_comment"] == "CHG1"

    # ikinci tarama: geçmiş ve durum kalıcılığı
    c.post("/api/scan", json={}, headers=H)
    c.wait_scan()
    assert len(c.get("/api/scan/history", headers=H).json()) == 2
    assert c.get("/api/summary", headers=H).json()["resolved_count"] == 1
    assert c.get("/api/summary", params={"session_id": sid}, headers=H).json()["session_id"] == sid

    # raporlar (token query ile, tarayıcıdan indirme gibi)
    tok = H["Authorization"][7:]
    for fmt in ("html", "xlsx", "csv", "json"):
        assert c.get(f"/api/report/{fmt}", params={"token": tok}).status_code == 200
    assert c.get("/api/report/pdf", params={"token": tok}).status_code == 400
    assert "WAF Healthcheck" in c.get("/").text


def test_resolved_change_detection(tmp_path):
    from f5waf.state import StateStore
    st = StateStore(str(tmp_path / "x.db"))
    f = {"fingerprint": "abc", "object_name": "/P/pol", "evidence": {"staged_signatures": 10, "observed_days": 3}}
    st.set_status("abc", "resolved", comment=None, assigned_to=None, user="u", finding=f)
    assert st.detect_changes([{**f, "evidence": {"staged_signatures": 10, "observed_days": 9}}]) == 0  # gün sayısı yok sayılır
    assert st.detect_changes([{**f, "evidence": {"staged_signatures": 25, "observed_days": 9}}]) == 1
    assert st.status_map()["abc"]["is_changed"] == 1


def test_current_view_merges_partial_scans(client):
    c = client
    H = c.H = _login(c)
    c.post("/api/scan", json={}, headers=H)
    c.wait_scan()
    full = c.get("/api/summary", headers=H).json()
    parts_full = {(p["device"], p["partition"]) for p in c.get("/api/partitions", headers=H).json()}
    # yalnızca tek müşteriyi yeniden tara
    c.post("/api/scan", json={"device_ids": ["waf01"], "partitions": "P_ACME"}, headers=H)
    c.wait_scan()
    cur = c.get("/api/summary", headers=H).json()
    assert cur["merged_from"] == 2
    assert cur["partitions"] == full["partitions"]               # diğer müşteriler ekrandan düşmedi
    assert cur["total_findings"] == full["total_findings"]       # mock deterministik: aynı sonuç
    assert {(p["device"], p["partition"]) for p in c.get("/api/partitions", headers=H).json()} == parts_full
    hist = c.get("/api/scan/history", headers=H).json()
    only = c.get("/api/summary", params={"session_id": hist[0]["id"]}, headers=H).json()
    assert only["partitions"] < full["partitions"]               # geçmiş görünümde tek tarama


def test_protections_inventory(client):
    c = client
    H = c.H = _login(c)
    c.post("/api/scan", json={}, headers=H)
    c.wait_scan()
    for kind in ("dos", "bot"):
        used = c.get("/api/protections", params={"kind": kind}, headers=H).json()
        unused = c.get("/api/protections", params={"kind": kind, "show": "unused"}, headers=H).json()
        assert used["stats"]["used"] == len(used["rows"]) > 0
        assert all(r[kind] for r in used["rows"])
        assert all(not r[kind] and r["eligible"] for r in unused["rows"])
        assert not any(r["virtual"].endswith("_80") for r in unused["rows"])   # redirect VS beklenmez
        m = next(iter(used["stats"]["by_mode"]))
        assert all(r[kind]["mode"] == m for r in c.get("/api/protections", params={"kind": kind, "mode": m},
                                                        headers=H).json()["rows"])
    assert c.get("/api/protections", params={"kind": "xyz"}, headers=H).status_code == 422
    # çoklu önem filtresi
    f = c.get("/api/findings", params={"severity": "CRITICAL,HIGH"}, headers=H).json()["findings"]
    assert f and {x["severity"] for x in f} <= {"CRITICAL", "HIGH"}
    # excel'de DoS/Bot sayfaları
    import io

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(c.get("/api/report/xlsx", headers=H).content))
    assert "DoS Protection" in wb.sheetnames and "Bot Protection" in wb.sheetnames


def test_dos_profile_parsing():
    from f5waf.protections import dos_info
    d = dos_info("/P/dos", {"applicationReference": {"items": [
        {"tpsBased": {"mode": "transparent"}, "stressBased": {"operationMode": "blocking"},
         "botDefense": {"mode": "during-attacks"}}]}})
    assert d["kind"] == "application" and d["mode"] == "blocking" and d["legacy_bot"] == "blocking"
    assert dos_info("/P/net", {})["kind"] == "network"


def test_policy_inventory(client):
    c = client
    H = c.H = _login(c)
    c.post("/api/scan", json={}, headers=H)
    c.wait_scan()
    d = c.get("/api/policies", headers=H).json()
    st = d["stats"]
    assert st["total"] == len(d["rows"]) > 0
    assert sum(st["by_mode"].values()) == st["total"]
    assert st["attached"] + st["orphan"] == st["total"]
    assert {x["device"] for x in st["by_device"]} == {"waf01", "waf02"}
    r = d["rows"][0]
    for k in ("device", "device_host", "partition", "customer", "full_path", "mode", "virtual_servers", "findings"):
        assert k in r
    # filtreler
    tr = c.get("/api/policies", params={"mode": "transparent"}, headers=H).json()["rows"]
    assert tr and all(x["mode"] == "transparent" for x in tr)
    assert any(x["transparent_days"] for x in tr)          # WAF-001 bulgusundan gün sayısı
    orphan = c.get("/api/policies", params={"vs": "orphan"}, headers=H).json()["rows"]
    assert all(not x["virtual_servers"] for x in orphan)
    w1 = c.get("/api/policies", params={"device": "waf01"}, headers=H).json()
    assert all(x["device"] == "waf01" for x in w1["rows"]) and len(w1["stats"]["by_device"]) == 2
    assert c.get("/api/policies", params={"search": "mock"}, headers=H).json()["rows"]   # host = config değeri
    vs = next(x for x in d["rows"] if x["virtual_servers"])["virtual_servers"][0]["name"]
    found = c.get("/api/policies", params={"search": vs.split("/")[-1]}, headers=H).json()["rows"]
    assert any(v["name"] == vs for x in found for v in x["virtual_servers"])
    import io

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(c.get("/api/report/xlsx", headers=H).content))
    assert wb["ASM Policyler"].max_row == st["total"] + 1


def test_dotenv_autoload_and_empty_credentials(tmp_path, monkeypatch):
    from f5waf.client import F5Client, F5Error
    from f5waf.models import Config
    monkeypatch.delenv("HC_TEST_USER", raising=False)
    monkeypatch.delenv("HC_TEST_PASS", raising=False)
    monkeypatch.delenv("HC_TEST_GONE", raising=False)
    (tmp_path / ".env").write_text('# yorum\nHC_TEST_USER=wafhc\nexport HC_TEST_PASS="p#ss=1"\n', encoding="utf-8")
    (tmp_path / "c.yaml").write_text(
        "defaults: {username: '${HC_TEST_USER}', password: '${HC_TEST_PASS}'}\n"
        "devices: [{name: a, host: 1.1.1.1}, {name: b, host: 1.1.1.2, password: '${HC_TEST_GONE}'}]\n",
        encoding="utf-8")
    cfg = Config.load(str(tmp_path / "c.yaml"))
    assert cfg.devices[0]["username"] == "wafhc" and cfg.devices[0]["password"] == "p#ss=1"
    assert cfg.missing_env == ["HC_TEST_GONE"]
    with pytest.raises(F5Error, match="boş"):
        F5Client("1.1.1.2", "wafhc", "").login()


def test_rejected_credentials_are_not_retried(monkeypatch):
    from f5waf import client as cmod
    cmod.clear_auth_failures()
    calls = []

    class R:
        status_code = 401
        text = '{"code":401,"message":"Maximum number of login attempts exceeded."}'

        def json(self):
            return {"code": 401, "message": "Maximum number of login attempts exceeded."}

    c = cmod.F5Client("10.0.0.1", "u", "p", verify_ssl=False)
    monkeypatch.setattr(c.session, "post", lambda *a, **k: calls.append(1) or R())
    with pytest.raises(cmod.F5Error, match="Unlock"):
        c.login()
    c2 = cmod.F5Client("10.0.0.1", "u", "p", verify_ssl=False)
    monkeypatch.setattr(c2.session, "post", lambda *a, **k: calls.append(1) or R())
    with pytest.raises(cmod.F5Error, match="tekrar denenmedi"):
        c2.login()
    assert len(calls) == 1          # ikinci tarama cihaza hiç gitmedi
    c3 = cmod.F5Client("10.0.0.1", "u", "yeni", verify_ssl=False)   # parola değişince denenir
    monkeypatch.setattr(c3.session, "post", lambda *a, **k: calls.append(1) or R())
    with pytest.raises(cmod.F5Error):
        c3.login()
    assert len(calls) == 2
    cmod.clear_auth_failures()
